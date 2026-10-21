"""真实本地存储/解析 API 集成测试；固定向量不评估模型召回质量。"""
import hashlib
import importlib.util
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch
from uuid import NAMESPACE_URL, uuid5

from research.documents import DocumentStore
from research.runtime import TaskBusy, TaskFileLock
from research.vector import HybridIndex, import_file


def available(name):
    return importlib.util.find_spec(name) is not None


class EmbeddingFixture:
    embedding_size = 2

    def __init__(self):
        self.batches = []

    def passage_embed(self, texts, **kwargs):
        self.batches.append(list(texts))
        return iter([[1., 0.] if "semantic" in text else [0., 1.]
                     for text in texts])

    def query_embed(self, texts, **kwargs):
        return iter([[1., 0.] for _ in texts])


class RerankerFixture:
    def rerank(self, query, texts, **kwargs):
        return iter([10. if "promote" in text else 0. for text in texts])


def write_text_pdf(path, pages):
    """仅用标准库生成带 Helvetica 文本的真实 PDF，不添加额外测试依赖。"""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    page_ids = []
    font_id = 3 + len(pages) * 2
    for text in pages:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append((
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {page_id + 1} 0 R >>"
        ).encode("ascii"))
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 20 200 Td ({escaped}) Tj ET".encode("ascii")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
                       + stream + b"\nendstream")
    kids = " ".join(f"{number} 0 R" for number in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode("ascii")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    output, offsets = bytearray(b"%PDF-1.4\n"), [0]
    for number, content in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii") + content + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode("ascii"))
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode("ascii")
                          for offset in offsets[1:]))
    output.extend((f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
                   f"startxref\n{xref}\n%%EOF\n").encode("ascii"))
    path.write_bytes(output)


class IntegrationBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.conn = sqlite3.connect(":memory:")
        self.addCleanup(self.conn.close)
        self.store = DocumentStore(self.conn)
        self.task = self.store.create_task("alpha research")
        self.store.update_task(self.task, "planned")
        # 本套测试禁止真实网络；真实 Qdrant 使用嵌入式本地模式。
        for target in ("requests.sessions.Session.request", "httpx.Client.send"):
            guard = patch(target, side_effect=AssertionError("network disabled"))
            guard.start()
            self.addCleanup(guard.stop)


class LocalParserIntegrationTests(IntegrationBase):
    def test_txt_and_markdown_metadata_and_idempotency(self):
        text = self.root / "资料 #1.txt"
        text.write_text("alpha 正文", encoding="utf-8-sig")
        markdown = self.root / "sections.md"
        markdown.write_text("# First\nalpha\n# Second\nbeta", encoding="utf-8")
        first = import_file(self.store, self.task, text)
        second = import_file(self.store, self.task, markdown)
        self.assertEqual(import_file(self.store, self.task, text), first)
        self.assertEqual(import_file(self.store, self.task, markdown), second)
        documents = self.store.list_documents(self.task)
        self.assertEqual(len(documents), 3)
        self.assertEqual([item["metadata"]["section"] for item in documents],
                         ["正文", "First", "Second"])
        self.assertEqual(documents[0]["metadata"]["sha256"],
                         hashlib.sha256(text.read_bytes()).hexdigest())
        self.assertIn("%23", documents[0]["url"])
        self.assertEqual(documents[0]["metadata"]["file_path"], str(text.resolve()))

    @unittest.skipUnless(available("pypdf"), "需要安装 pypdf 才能验证真实 PDF 解析")
    def test_real_pdf_page_text_and_page_provenance(self):
        path = self.root / "two-pages.pdf"
        write_text_pdf(path, ["alpha on first page", "beta on second page"])
        import_file(self.store, self.task, path)
        documents = self.store.list_documents(self.task)
        self.assertEqual([item["metadata"]["page"] for item in documents], [1, 2])
        self.assertEqual([item["metadata"]["section"] for item in documents],
                         ["第 1 页", "第 2 页"])
        self.assertEqual([item["content"] for item in self.store.list_chunks(self.task)],
                         ["alpha on first page", "beta on second page"])

    @unittest.skipUnless(available("pypdf"), "需要安装 pypdf 才能验证加密 PDF 拒绝")
    def test_encrypted_pdf_is_rejected_without_partial_document(self):
        from pypdf import PdfWriter
        path = self.root / "encrypted.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.encrypt("secret")
        with path.open("wb") as output:
            writer.write(output)
        with self.assertRaisesRegex(ValueError, "加密 PDF"):
            import_file(self.store, self.task, path)
        self.assertEqual(self.store.list_documents(self.task), [])

    @unittest.skipUnless(available("docx"), "需要安装 python-docx 才能验证真实 DOCX 解析")
    def test_real_docx_preserves_headings_and_table_order(self):
        from docx import Document
        path = self.root / "structured.docx"
        document = Document()
        document.add_heading("First", level=1)
        document.add_paragraph("alpha before table")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
        document.add_paragraph("alpha after table")
        document.add_heading("Second", level=1)
        document.add_paragraph("beta")
        document.save(path)
        import_file(self.store, self.task, path)
        documents = self.store.list_documents(self.task)
        self.assertEqual([item["metadata"]["section"] for item in documents],
                         ["First", "Second"])
        self.assertEqual(self.store.list_chunks(self.task)[0]["content"],
                         "alpha before table\nName | Value\nalpha after table")
        self.assertEqual(import_file(self.store, self.task, path),
                         [item["id"] for item in documents])


@unittest.skipUnless(available("qdrant_client"), "需要安装 qdrant-client 才能验证真实本地索引")
class QdrantIntegrationTests(IntegrationBase):
    def setUp(self):
        super().setUp()
        self.embedder = EmbeddingFixture()
        self.index = HybridIndex(self.root / "qdrant", embedder=self.embedder,
                                 reranker=RerankerFixture())
        self.chunks = [
            {"evidence_id": "D1C1", "document_id": 1, "chunk_index": 1,
             "title": "Alpha", "content": "alpha lexical", "source_type": "web_text"},
            {"evidence_id": "D2C1", "document_id": 2, "chunk_index": 1,
             "title": "Beta", "content": "semantic promote", "source_type": "web_text"},
        ]

    def test_real_qdrant_persistence_reopen_and_incremental_update(self):
        from qdrant_client import QdrantClient
        result = self.index.search(self.task, "alpha", self.chunks)
        self.assertEqual([item["evidence_id"] for item in result], ["D2C1", "D1C1"])
        reopened = HybridIndex(self.index.path, embedder=self.embedder,
                               reranker=RerankerFixture())
        self.assertEqual(reopened.search(self.task, "alpha", self.chunks), result)
        self.assertEqual([len(batch) for batch in self.embedder.batches], [2])
        self.chunks[0]["content"] += " changed"
        reopened.search(self.task, "alpha", self.chunks)
        self.assertEqual([len(batch) for batch in self.embedder.batches], [2, 1])
        with closing(QdrantClient(path=str(self.index.path))) as client:
            self.assertEqual(client.count(self.index._collection(), exact=True).count, 2)

    def test_task_and_allowed_point_filters_use_real_qdrant_models(self):
        from qdrant_client import QdrantClient, models
        self.index.search("other", "alpha", self.chunks)
        self.index.search(self.task, "alpha", self.chunks)
        with closing(QdrantClient(path=str(self.index.path))) as client:
            client.upsert(self.index._collection(), points=[models.PointStruct(
                id=str(uuid5(NAMESPACE_URL, self.task + "/D999C1")),
                vector=[1., 0.], payload={"task_id": self.task, "content": "SECRET"},
            )], wait=True)
        dense = self.index._dense(self.task, "alpha", self.chunks)
        self.assertEqual({item["evidence_id"] for item in dense}, {"D1C1", "D2C1"})
        self.assertNotIn("SECRET", str(dense))
        self.assertEqual(len(self.embedder.batches), 2)

    def test_revision_change_uses_separate_real_collection(self):
        from qdrant_client import QdrantClient
        self.index.search(self.task, "alpha", self.chunks)
        changed = HybridIndex(self.index.path, revision="new-revision",
                              embedder=self.embedder, reranker=RerankerFixture())
        changed.search(self.task, "alpha", self.chunks)
        self.assertNotEqual(self.index._collection(), changed._collection())
        with closing(QdrantClient(path=str(self.index.path))) as client:
            for name in (self.index._collection(), changed._collection()):
                self.assertEqual(client.count(name, exact=True).count, 2)

    def test_global_index_lock_is_released_after_busy_attempt(self):
        lock_path = self.index.path.parent / (self.index.path.name + ".db")
        with TaskFileLock(lock_path, "qdrant-global"):
            with self.assertRaises(TaskBusy):
                self.index.search(self.task, "alpha", self.chunks)
        self.assertEqual(len(self.index.search(self.task, "alpha", self.chunks)), 2)


if __name__ == "__main__":
    unittest.main()
